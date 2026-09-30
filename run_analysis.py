# Generate all project outputs
import os, sys
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.ticker as mtick
import seaborn as sns
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

DATA_PATH = Path('data/pnl_data.xlsx')
OUT_CHARTS = Path('outputs/charts')
OUT_CHARTS.mkdir(parents=True, exist_ok=True)
OUT_REPORTS = Path('outputs/reports')
OUT_REPORTS.mkdir(parents=True, exist_ok=True)

# 1. Load Raw Data
raw = pd.read_excel(DATA_PATH, sheet_name='Raw Data', nrows=13)
raw = raw.dropna(subset=['Month'])
raw = raw[raw['Month'] != 'Note:'].reset_index(drop=True)
raw.columns = raw.columns.str.strip()

# 2. Build P&L
TAX_RATE = 0.20
df = raw.copy()
df['Gross_Profit']   = df['Revenue'] - df['COGS']
df['Gross_Margin']   = df['Gross_Profit'] / df['Revenue']
df['Total_OpEx']     = df[['Salaries','Marketing','Rent','Utilities','Other Opex']].sum(axis=1)
df['OpEx_pct_Rev']   = df['Total_OpEx'] / df['Revenue']
df['EBITDA']         = df['Gross_Profit'] - df['Total_OpEx']
df['EBITDA_Margin']  = df['EBITDA'] / df['Revenue']
df['EBIT']           = df['EBITDA'] - df['Depreciation']
df['PBT']            = df['EBIT'] - df['Interest']
df['Tax_Calc']       = (df['PBT'] * TAX_RATE).round(2)
df['Net_Profit']     = df['PBT'] - df['Tax_Calc']
df['NP_Margin']      = df['Net_Profit'] / df['Revenue']
df['Rev_MoM_Growth'] = df['Revenue'].pct_change()

MONTH_ORDER = ['April','May','June','July','August','September',
               'October','November','December','January','February','March']
df['Month'] = pd.Categorical(df['Month'], categories=MONTH_ORDER, ordered=True)
df = df.sort_values('Month').reset_index(drop=True)

fy_revenue     = df['Revenue'].sum()
fy_gp          = df['Gross_Profit'].sum()
fy_gm          = fy_gp / fy_revenue
fy_ebitda      = df['EBITDA'].sum()
fy_ebitda_m    = fy_ebitda / fy_revenue
fy_np          = df['Net_Profit'].sum()
fy_np_m        = fy_np / fy_revenue
fy_opex        = df['Total_OpEx'].sum()
fy_opex_pct    = fy_opex / fy_revenue

months_short = ['Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec','Jan','Feb','Mar']
DARK_BLUE  = '#17365D'
TEAL       = '#2F7E79'
AMBER      = '#C55A11'
PURPLE     = '#7030A0'
RED        = '#C00000'
GREEN_OK   = '#276221'

# Chart 1: KPI Cards
fig, axes = plt.subplots(1, 4, figsize=(16, 3.5))
fig.patch.set_facecolor('#17365D')
kpis = [
    ('Total Revenue', f'Rs {fy_revenue:,.0f}L', 'Full Year Total', TEAL),
    ('Gross Margin', f'{fy_gm:.1%}', 'Full Year Avg', AMBER),
    ('EBITDA', f'Rs {fy_ebitda:,.1f}L', 'Full Year Total', PURPLE),
    ('Net Profit', f'Rs {fy_np:,.1f}L', 'Full Year Total', '#E8A020'),
]
for ax, (title, value, sub, color) in zip(axes, kpis):
    ax.set_facecolor('#1E4080')
    ax.set_xlim(0,1); ax.set_ylim(0,1)
    ax.axis('off')
    ax.text(0.5, 0.72, title, ha='center', va='center', color='#A9C4E8', fontsize=11, fontweight='bold', transform=ax.transAxes)
    ax.text(0.5, 0.42, value, ha='center', va='center', color='white', fontsize=24, fontweight='bold', transform=ax.transAxes)
    ax.text(0.5, 0.16, sub, ha='center', va='center', color=color, fontsize=10, transform=ax.transAxes)
    for spine in ax.spines.values():
        spine.set_edgecolor(color); spine.set_linewidth(2); spine.set_visible(True)
fig.suptitle('FY Financial Performance Dashboard -- Rs Lakhs', color='white', fontsize=14, fontweight='bold', y=1.02)
plt.tight_layout(pad=0.8)
plt.savefig(OUT_CHARTS / 'kpi_cards.png', bbox_inches='tight', facecolor='#17365D', dpi=150)
plt.close()

# Chart 2: Revenue Trend & MoM
cagr_mom = (df['Revenue'].iloc[-1] / df['Revenue'].iloc[0]) ** (1/11) - 1
fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(14, 8), sharex=True, gridspec_kw={'height_ratios':[3,1.5]})
bars = ax1.bar(months_short, df['Revenue'], color=DARK_BLUE, alpha=0.85, width=0.6, zorder=3, label='Revenue (Rs L)')
ax1.plot(months_short, df['Revenue'], color=TEAL, marker='o', lw=2.5, markersize=6, zorder=4)
for bar in bars:
    ax1.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 0.8, f'{bar.get_height():.0f}L', ha='center', va='bottom', fontsize=9, color=DARK_BLUE, fontweight='bold')
mean_rev = df['Revenue'].mean()
ax1.axhline(mean_rev, color=AMBER, ls='--', lw=1.5, label=f'Monthly Avg Rs {mean_rev:.0f}L')
ax1.set_ylabel('Revenue (Rs Lakhs)')
ax1.set_title('Monthly Revenue Trend')
ax1.legend(loc='upper left')
ax1.set_ylim(60, 135)

colors_mom = [GREEN_OK if x >= 0 else RED for x in df['Rev_MoM_Growth'].fillna(0)]
ax2.bar(months_short, df['Rev_MoM_Growth'].fillna(0) * 100, color=colors_mom, alpha=0.8, width=0.6, zorder=3)
ax2.axhline(0, color='black', lw=0.8)
ax2.axhline(cagr_mom * 100, color=AMBER, ls='--', lw=1.5, label=f'Compound MoM avg {cagr_mom:.1%}')
ax2.set_ylabel('MoM Growth (%)')
ax2.set_title('Month-on-Month Revenue Growth')
ax2.legend(loc='lower right')
ax2.yaxis.set_major_formatter(mtick.FormatStrFormatter('%.0f%%'))
plt.tight_layout()
plt.savefig(OUT_CHARTS / 'revenue_trend.png', bbox_inches='tight', dpi=150)
plt.close()

# Chart 3: Margin Trends
fig, ax = plt.subplots(figsize=(14, 6))
ax.plot(months_short, df['Gross_Margin']*100, marker='o', color=DARK_BLUE, lw=2.5, markersize=7, label='Gross Margin %')
ax.plot(months_short, df['EBITDA_Margin']*100, marker='s', color=TEAL, lw=2.5, markersize=7, label='EBITDA Margin %')
ax.plot(months_short, df['NP_Margin']*100, marker='^', color=AMBER, lw=2.5, markersize=7, label='Net Profit Margin %')
ax.plot(months_short, df['OpEx_pct_Rev']*100, marker='D', color=RED, lw=2, markersize=6, ls='--', label='OpEx as % of Revenue')
ax.set_ylabel('Margin (%)')
ax.set_title('Monthly Margin Trends -- Gross / EBITDA / Net Profit vs OpEx %')
ax.legend(loc='upper left', fontsize=10)
ax.yaxis.set_major_formatter(mtick.FormatStrFormatter('%.0f%%'))
plt.tight_layout()
plt.savefig(OUT_CHARTS / 'margin_trends.png', bbox_inches='tight', dpi=150)
plt.close()

# Chart 4: Budget vs Actual
budget = {'Revenue': 1150, 'COGS': 630, 'Gross Profit': 520, 'Operating Expenses': 275, 'EBITDA': 245, 'Net Profit': 145}
actual = {'Revenue': df['Revenue'].sum(), 'COGS': df['COGS'].sum(), 'Gross Profit': df['Gross_Profit'].sum(), 'Operating Expenses': df['Total_OpEx'].sum(), 'EBITDA': df['EBITDA'].sum(), 'Net Profit': df['Net_Profit'].sum()}
fav_if_pos = {'Revenue': True, 'COGS': False, 'Gross Profit': True, 'Operating Expenses': False, 'EBITDA': True, 'Net Profit': True}
metrics = list(budget.keys())
b_vals = [budget[m] for m in metrics]
a_vals = [round(actual[m], 1) for m in metrics]
x = np.arange(len(metrics))
width = 0.36
fig, ax = plt.subplots(figsize=(14, 6))
bars_b = ax.bar(x - width/2, b_vals, width, label='Budget', color=DARK_BLUE, alpha=0.85)
bars_a = ax.bar(x + width/2, a_vals, width, label='Actual', color=[(GREEN_OK if (actual[m]-budget[m]>0 and fav_if_pos[m]) or (actual[m]-budget[m]<0 and not fav_if_pos[m]) else RED) for m in metrics], alpha=0.85)
for i, m in enumerate(metrics):
    var = actual[m] - budget[m]
    is_fav = (var > 0 and fav_if_pos[m]) or (var < 0 and not fav_if_pos[m])
    col = GREEN_OK if is_fav else RED
    label = f'+{var:.1f}' if var >= 0 else f'{var:.1f}'
    ax.text(x[i] + width/2, max(b_vals[i], a_vals[i]) + 5, label, ha='center', fontsize=9, color=col, fontweight='bold')
ax.set_xticks(x)
ax.set_xticklabels(metrics, rotation=15, ha='right')
ax.set_ylabel('Rs Lakhs')
ax.set_title('Budget vs Actual -- Full Year (Rs Lakhs)')
ax.legend()
plt.tight_layout()
plt.savefig(OUT_CHARTS / 'budget_vs_actual.png', bbox_inches='tight', dpi=150)
plt.close()

# Chart 5: OpEx Analysis
opex_cols = ['Salaries','Marketing','Rent','Utilities','Other Opex']
fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 6))
bottom = np.zeros(12)
for col, color in zip(opex_cols, [DARK_BLUE, TEAL, AMBER, PURPLE, '#E8A020']):
    ax1.bar(months_short, df[col], bottom=bottom, label=col, color=color, alpha=0.85)
    bottom += df[col].values
ax1.plot(months_short, df['Total_OpEx'], color='black', marker='o', lw=2, markersize=5, label='Total OpEx', zorder=5)
ax1.set_title('OpEx Composition by Month')
ax1.set_ylabel('Rs Lakhs')
ax1.legend(fontsize=9, loc='upper left')

budget_opex_pct = 275 / 1150
ax2.plot(months_short, df['OpEx_pct_Rev']*100, marker='o', color=RED, lw=2.5, markersize=7, label='Actual OpEx % Revenue')
ax2.axhline(budget_opex_pct*100, color=DARK_BLUE, ls='--', lw=2, label=f'Budget benchmark {budget_opex_pct:.1%}')
ax2.fill_between(months_short, df['OpEx_pct_Rev']*100, budget_opex_pct*100, where=(df['OpEx_pct_Rev'] > budget_opex_pct), alpha=0.15, color=RED, label='Over budget zone')
ax2.set_title('OpEx % of Revenue vs Budget Benchmark')
ax2.yaxis.set_major_formatter(mtick.FormatStrFormatter('%.0f%%'))
ax2.legend(fontsize=9)
plt.tight_layout()
plt.savefig(OUT_CHARTS / 'opex_analysis.png', bbox_inches='tight', dpi=150)
plt.close()

# Chart 6: Tax Sensitivity
rates = [0.15, 0.18, 0.20, 0.22, 0.25, 0.30]
np_vals = [(df['PBT'] - (df['PBT'] * r).round(2)).sum() for r in rates]
fig, ax = plt.subplots(figsize=(10, 5))
bar_colors = [GREEN_OK if v >= df['Net_Profit'].sum() else RED for v in np_vals]
ax.bar([f'{r:.0%}' for r in rates], np_vals, color=bar_colors, alpha=0.85)
ax.axhline(df['Net_Profit'].sum(), color=DARK_BLUE, ls='--', lw=1.5, label=f'Current NP @ {TAX_RATE:.0%}: Rs {df["Net_Profit"].sum():.1f}L')
for i, v in enumerate(np_vals):
    ax.text(i, v + 0.5, f'Rs {v:.1f}L', ha='center', fontsize=9, fontweight='bold')
ax.set_xlabel('Corporate Tax Rate Assumption')
ax.set_ylabel('Full-Year Net Profit (Rs Lakhs)')
ax.set_title('Tax Rate Sensitivity -- Impact on Net Profit')
ax.legend()
plt.tight_layout()
plt.savefig(OUT_CHARTS / 'tax_sensitivity.png', bbox_inches='tight', dpi=150)
plt.close()

# Chart 7: Correlation Heatmap
corr_cols = ['Revenue','Gross_Profit','Total_OpEx','EBITDA','Net_Profit','Gross_Margin','EBITDA_Margin','NP_Margin','OpEx_pct_Rev']
corr = df[corr_cols].corr()
fig, ax = plt.subplots(figsize=(10, 8))
sns.heatmap(corr, annot=True, fmt='.2f', cmap='coolwarm', vmin=-1, vmax=1, square=True, linewidths=0.5, ax=ax, annot_kws={'size': 9})
ax.set_title('Correlation Matrix -- Key Financial Metrics')
plt.tight_layout()
plt.savefig(OUT_CHARTS / 'correlation_heatmap.png', bbox_inches='tight', dpi=150)
plt.close()

# 3. Excel Export Report
wb_out = Workbook()
ws_fy = wb_out.active
ws_fy.title = 'Full-Year Summary'
header_fill = PatternFill('solid', fgColor='17365D')
header_font = Font(color='FFFFFF', bold=True)
ws_fy.append(['Metric', 'Value', 'Unit'])
for c in ws_fy[1]:
    c.font = header_font; c.fill = header_fill

for row in [
    ('Revenue', fy_revenue, 'Rs Lakhs'),
    ('Gross Profit', fy_gp, 'Rs Lakhs'),
    ('Gross Margin', f'{fy_gm:.1%}', '%'),
    ('EBITDA', fy_ebitda, 'Rs Lakhs'),
    ('EBITDA Margin', f'{fy_ebitda_m:.1%}', '%'),
    ('Net Profit', fy_np, 'Rs Lakhs'),
    ('Net Profit Margin', f'{fy_np_m:.1%}', '%'),
    ('OpEx % Revenue', f'{fy_opex_pct:.1%}', '%'),
    ('Tax Rate Assumed', f'{TAX_RATE:.0%}', '%'),
]:
    ws_fy.append(list(row))

wb_out.save(OUT_REPORTS / 'PnL_Summary_Report.xlsx')
print('SUCCESS: All 7 charts and Excel summary report created!')
